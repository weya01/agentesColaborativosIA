#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Converter RELATORIO_TECNICO.md para .docx
"""

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import os
    
    # Ler o ficheiro Markdown
    with open('docs/RELATORIO_TECNICO.md', 'r', encoding='utf-8') as f:
        conteudo = f.read()
    
    # Criar documento Word
    doc = Document()
    
    # Adicionar título
    title = doc.add_heading('RELATÓRIO TÉCNICO', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Sistema de Simulação Multi-Agente Colaborativo')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    
    # Processar linhas do conteúdo
    linhas = conteudo.split('\n')
    
    for linha in linhas:
        if not linha.strip():
            doc.add_paragraph('')
        elif linha.startswith('# '):
            doc.add_heading(linha[2:], 1)
        elif linha.startswith('## '):
            doc.add_heading(linha[3:], 2)
        elif linha.startswith('### '):
            doc.add_heading(linha[4:], 3)
        elif linha.startswith('#### '):
            doc.add_heading(linha[5:], 4)
        elif linha.startswith('- '):
            doc.add_paragraph(linha[2:], style='List Bullet')
        elif linha.startswith('| '):
            # Pular tabelas complexas por enquanto
            continue
        else:
            doc.add_paragraph(linha)
    
    # Guardar documento
    output_path = 'docs/Relatorio_IA2024_GRUPO_26_TECNICO.docx'
    doc.save(output_path)
    
    print(f'[OK] Documento Word criado com sucesso: {output_path}')
    print(f'[OK] Tamanho: {os.path.getsize(output_path) / 1024:.1f} KB')
    
except ImportError:
    print('python-docx não está instalado. Instalando...')
    import subprocess
    subprocess.check_call(['pip', 'install', 'python-docx', '-q'])
    print('Instalação concluída. Execute novamente o script.')
except Exception as e:
    print(f'Erro: {e}')
